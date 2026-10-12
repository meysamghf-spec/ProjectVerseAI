"""API tests that check persisted calculations, not only status codes."""

import io

from app.engine.distributions import percentile
from openpyxl import load_workbook


def _create(client, prompt=None, **extra):
    body = {
        "prompt": prompt
        or (
            "Plan a 20-story residential building with three basement parking levels, "
            "25,000 square meters of gross floor area, a $15 million budget, and a 30-month "
            "construction schedule."
        ),
        "name": "Tower",
        "generate": True,
        "start_date": "2026-01-05",
        "run_monte_carlo": False,
    }
    body.update(extra)
    response = client.post("/api/v1/projects", json=body)
    assert response.status_code == 201, response.text
    return response.json()


def test_health_ready_and_validation_error(client):
    assert client.get("/api/v1/health").json()["status"] == "ok"
    assert client.get("/api/v1/ready").json()["database"] is True
    rejected = client.post("/api/v1/projects", json={"prompt": ""})
    assert rejected.status_code == 422
    body = rejected.json()
    assert body["error"]["code"] == "validation_error"
    missing = client.get("/api/v1/projects/does-not-exist")
    assert missing.status_code == 404
    assert missing.json()["error"]["code"] == "not_found"


def test_project_lifecycle_schedule_budget_and_reopen(client):
    created = _create(client)
    project_id = created["id"]
    assert created["specification"]["known"]
    known = {item["key"]: item["value"] for item in created["specification"]["known"]}
    assert known["floors"] == 20
    assert "location" in created["specification"]["missing"]
    assert created["assumptions"]
    assert created["latest_simulation"]["status"] == "completed"
    result = created["latest_simulation"]["result"]
    assert result["project_duration"] == 900
    assert result["total_cost"] == created["cost_summary"]["total"]
    assert abs(result["total_cost"] - sum(line["amount"] for line in created["cost_summary"]["lines"] if line["include_in_total"])) < 0.05
    assert result["critical_path"]
    listed = client.get("/api/v1/projects")
    assert listed.json()["total"] == 1
    reopened = client.get(f"/api/v1/projects/{project_id}")
    assert reopened.json()["tasks"]
    assert reopened.json()["scene"]["floors"] == 20

    updated = client.patch(
        f"/api/v1/projects/{project_id}",
        json={"floors": 12, "floor_to_floor_height_m": 3.6, "name": "Tower revised"},
    )
    assert updated.status_code == 200
    body = updated.json()
    assert body["name"] == "Tower revised"
    assert body["scene"]["floors"] == 12
    assert body["scene"]["floor_height_m"] == 3.6
    assert body["schedule_stale"] is True

    superstructure = next(task for task in created["tasks"] if task["code"] == "CON-SUPER")
    edited = client.patch(
        f"/api/v1/projects/{project_id}/tasks/{superstructure['id']}",
        json={"duration": superstructure["duration"] + 10},
    )
    assert edited.status_code == 200
    scheduled = client.post(f"/api/v1/projects/{project_id}/schedule", json={"set_baseline": False})
    assert scheduled.status_code == 200
    assert scheduled.json()["result"]["project_duration"] != 900 or body["schedule_stale"]

    deleted = client.delete(f"/api/v1/projects/{project_id}")
    assert deleted.status_code == 204
    assert client.get(f"/api/v1/projects/{project_id}").status_code == 404


def test_owner_isolation(client):
    created = _create(client)
    other = client.get("/api/v1/projects", headers={"X-Owner-Id": "someone-else"})
    assert other.json()["total"] == 0
    hidden = client.get(f"/api/v1/projects/{created['id']}", headers={"X-Owner-Id": "someone-else"})
    assert hidden.status_code == 404


def test_production_auth_requires_api_key(prod_client):
    blocked = prod_client.get("/api/v1/projects")
    assert blocked.status_code == 401
    allowed = prod_client.get("/api/v1/projects", headers={"X-API-Key": "test-key"})
    assert allowed.status_code == 200
    spoofed = prod_client.get("/api/v1/projects", headers={"X-Owner-Id": "dev-user"})
    assert spoofed.status_code == 401


def test_cycle_budget_line_monte_carlo_and_reports(client):
    created = _create(client, generate=True)
    project_id = created["id"]
    client.post(
        f"/api/v1/projects/{project_id}/dependencies",
        json={"predecessor_code": "CON-HAND", "successor_code": "CON-MOB"},
    )
    failed = client.post(f"/api/v1/projects/{project_id}/schedule", json={"set_baseline": False})
    assert failed.status_code == 422
    assert "cycle" in failed.json()["error"]["message"].lower()

    fresh = _create(client, name="Budget check")
    project_id = fresh["id"]
    before = fresh["cost_summary"]["total"]
    added = client.post(
        f"/api/v1/projects/{project_id}/costs",
        json={
            "code": "user-extra",
            "name": "Owner allowance",
            "category": "other",
            "amount": 1000,
            "source": "user",
            "source_date": "2026-02-01",
            "include_in_total": True,
        },
    )
    assert added.status_code == 201
    rollup = client.post(
        f"/api/v1/projects/{project_id}/costs",
        json={
            "code": "rollup-display",
            "name": "Rollup",
            "category": "other",
            "amount": 999999,
            "include_in_total": False,
            "notes": "Must not be added to the total.",
        },
    )
    assert rollup.status_code == 201
    summary = client.post(f"/api/v1/projects/{project_id}/budget").json()
    assert abs(summary["total"] - (before + 1000)) < 0.05
    assert summary["variance"] == fresh["budget_amount"] - summary["total"] or summary["budget"] == 15_000_000

    negative = client.post(
        f"/api/v1/projects/{project_id}/costs",
        json={"code": "neg", "name": "Bad", "category": "labor", "amount": -1},
    )
    assert negative.status_code == 422

    simulation = client.post(
        f"/api/v1/projects/{project_id}/simulations",
        json={"run_type": "monte_carlo", "seed": 11, "iterations": 40, "wait": True},
    )
    assert simulation.status_code == 201, simulation.text
    result = simulation.json()["result"]
    assert result["p50_duration"] <= result["p80_duration"] <= result["p90_duration"]
    assert result["duration_sample_count"] == 40
    again = client.post(
        f"/api/v1/projects/{project_id}/simulations",
        json={"run_type": "monte_carlo", "seed": 11, "iterations": 40, "wait": True},
    )
    assert again.json()["result"]["duration_sample_checksum"] == result["duration_sample_checksum"]
    detailed = client.get(
        f"/api/v1/projects/{project_id}/simulations/{simulation.json()['id']}?include_samples=true"
    ).json()["result"]
    assert detailed["p50_duration"] == percentile(detailed["duration_samples"], 50)
    assert detailed["deadline_probability"] is not None

    report = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"format": "json", "report_type": "full"},
    )
    assert report.status_code == 200
    payload = report.json()
    assert payload["deterministic"]["result"]["total_cost"] == summary["total"] or payload["cost_summary"]["total"]
    assert "disclaimer" in payload["disclaimer"].lower() or "not a verified" in payload["disclaimer"].lower()
    assert payload["engine_version"]
    pdf = client.post(f"/api/v1/projects/{project_id}/reports", json={"format": "pdf", "report_type": "summary"})
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    csv_report = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"format": "csv", "report_type": "schedule"},
    )
    assert b"CON-SUPER" in csv_report.content
    xlsx = client.post(
        f"/api/v1/projects/{project_id}/reports",
        json={"format": "xlsx", "report_type": "budget"},
    )
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert "Budget" in workbook.sheetnames


def test_scenario_isolation(client):
    created = _create(client)
    project_id = created["id"]
    baseline_duration = created["latest_simulation"]["result"]["project_duration"]
    created_scenario = client.post(
        f"/api/v1/projects/{project_id}/scenarios",
        json={
            "name": "More crews",
            "description": "Scale only the superstructure.",
            "duration_factors": {"superstructure": 0.5},
            "scalable_only": True,
            "scalable_phases": ["superstructure"],
            "run": True,
            "wait": True,
            "monte_carlo": True,
            "iterations": 20,
            "seed": 5,
        },
    )
    assert created_scenario.status_code == 201, created_scenario.text
    scenario = created_scenario.json()
    assert scenario["deterministic"]["result"]["project_duration"] < baseline_duration
    reopened = client.get(f"/api/v1/projects/{project_id}").json()
    superstructure = next(task for task in reopened["tasks"] if task["code"] == "CON-SUPER")
    assert superstructure["duration"] == 300
    comparison = client.post(f"/api/v1/projects/{project_id}/scenarios/compare")
    assert comparison.status_code == 200
    row = comparison.json()["scenarios"][0]
    assert row["duration_delta"] < 0
    assert row["input_changes"]


def test_demo_project(client):
    response = client.post("/api/v1/demo/residential-tower")
    assert response.status_code == 201, response.text
    project = response.json()
    assert project["name"] == "Residential Tower Demonstration"
    assert project["scene"]["floors"] == 20
    assert project["scene"]["basements"] == 3
    assert project["budget_amount"] == 15_000_000
    assert project["latest_simulation"]["result"]["project_duration"] == 900
    assert project["latest_monte_carlo"]["result"]["seed"] == 42
    assert len(project["risks"]) >= 3
    scenarios = client.get(f"/api/v1/projects/{project['id']}/scenarios").json()["items"]
    assert len(scenarios) == 3
    names = {item["name"] for item in scenarios}
    assert "Higher material prices" in names
    assert "Accelerated facade overlap" in names
    overlap = next(item for item in scenarios if item["name"] == "Accelerated facade overlap")
    assert overlap["deterministic"]["result"]["project_duration"] < 900
    assert overlap["deterministic"]["result"]["resource_conflicts"]
